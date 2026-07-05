import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser

from accounts.permissions import IsAdmin, IsActiveStudent, IsActiveLecturer, IsPGCoordinator
from courses.models import Enrollment, CourseAllocation
from notifications.utils import notify
from .models import SemesterResultBatch, Result, SemesterGPA, CumulativeGPA
from .signals import recompute_gpa_for_batch
from .serializers import (
    ResultSerializer, ResultUploadSerializer, SemesterResultBatchSerializer,
    SemesterGPASerializer, CumulativeGPASerializer,
)

SCORE_SHEET_COLUMNS = ["#", "Registration No.", "Name", "CA", "Exam"]


class StudentResultsView(APIView):
    """GET /api/results/mine/?session=<id>&semester=<FIRST|SECOND>"""
    permission_classes = [IsActiveStudent]

    def get(self, request):
        qs = Result.objects.filter(
            enrollment__student=request.user.student_profile,
            batch__status=SemesterResultBatch.Status.APPROVED,
        ).select_related(
            "enrollment__allocation__course", "batch__session"
        )
        session_id = request.query_params.get("session")
        semester = request.query_params.get("semester")
        if session_id:
            qs = qs.filter(batch__session_id=session_id)
        if semester:
            qs = qs.filter(batch__semester=semester)
        return Response(ResultSerializer(qs, many=True).data)


class StudentGPAView(APIView):
    """GET /api/results/gpa/  — returns all semester GPAs + cumulative."""
    permission_classes = [IsActiveStudent]

    def get(self, request):
        profile = request.user.student_profile
        semester_gpas = SemesterGPA.objects.filter(
            student=profile
        ).select_related("session").order_by("session__start_date", "semester")
        try:
            cgpa = CumulativeGPA.objects.get(student=profile)
            cgpa_data = CumulativeGPASerializer(cgpa).data
        except CumulativeGPA.DoesNotExist:
            cgpa_data = None
        return Response({
            "semester_gpas": SemesterGPASerializer(semester_gpas, many=True).data,
            "cumulative": cgpa_data,
        })


class LecturerResultUploadView(APIView):
    """POST /api/results/upload/  — lecturer uploads results for a batch."""
    permission_classes = [IsActiveLecturer]

    def post(self, request):
        batch_id = request.data.get("batch")
        results_data = request.data.get("results", [])

        try:
            batch = SemesterResultBatch.objects.get(pk=batch_id)
        except SemesterResultBatch.DoesNotExist:
            return Response({"detail": "Batch not found."}, status=404)

        if batch.status == SemesterResultBatch.Status.APPROVED:
            return Response({"detail": "Cannot modify an approved batch."}, status=400)

        errors = []
        created = 0
        with transaction.atomic():
            for item in results_data:
                s = ResultUploadSerializer(data=item)
                if not s.is_valid():
                    errors.append({"enrollment": item.get("enrollment"), "errors": s.errors})
                    continue
                d = s.validated_data
                try:
                    enrollment = Enrollment.objects.get(pk=d["enrollment"])
                except Enrollment.DoesNotExist:
                    errors.append({"enrollment": d["enrollment"], "errors": "Enrollment not found."})
                    continue
                Result.objects.update_or_create(
                    enrollment=enrollment,
                    defaults={"score": d["score"], "batch": batch}
                )
                created += 1

        return Response({"uploaded": created, "errors": errors}, status=200)


class ScoreSheetTemplateView(APIView):
    """
    GET /api/results/template/<allocation_pk>/
    Downloads a blank score sheet (xlsx) for the lecturer's allocation,
    pre-filled with the approved students; CA and Exam columns left empty.
    Layout mirrors the departmental score sheet format
    (e.g. CSC3315_2023_2024_SCORE_SHEET.xlsx).
    """
    permission_classes = [IsActiveLecturer]

    def get(self, request, allocation_pk):
        try:
            allocation = CourseAllocation.objects.select_related(
                "course__programme", "session"
            ).get(pk=allocation_pk, lecturer=request.user.lecturer_profile)
        except CourseAllocation.DoesNotExist:
            return Response({"detail": "Allocation not found."}, status=404)

        enrollments = Enrollment.objects.filter(
            allocation=allocation, status="APPROVED"
        ).select_related("student__user").order_by("student__matric_number")

        course = allocation.course
        session_name = allocation.session.name
        semester_display = allocation.get_semester_display()

        wb = Workbook()
        ws = wb.active
        ws.title = "Score Sheet"

        header_lines = [
            "Aliko Dangote University of Science and Technology, Wudil",
            "Faculty of Computing and Mathematical Science",
            "Department of Computer Science",
            f"{course.programme.get_programme_type_display()} - Score Sheet",
            f"{course.code} : {course.title}",
            f"Session: {session_name},          Semester: {semester_display},          "
            f"Date: {timezone.localdate().strftime('%d/%m/%Y')}",
            f"Number of Students: {enrollments.count()}",
            course.code,
        ]
        for i, line in enumerate(header_lines, start=1):
            ws.cell(row=i, column=1, value=line)
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
            cell = ws.cell(row=i, column=1)
            cell.font = Font(bold=(i <= 5))
            cell.alignment = Alignment(horizontal="center")

        header_row = len(header_lines) + 1
        for col, name in enumerate(SCORE_SHEET_COLUMNS, start=1):
            c = ws.cell(row=header_row, column=col, value=name)
            c.font = Font(bold=True)

        for idx, e in enumerate(enrollments, start=1):
            row = header_row + idx
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=e.student.matric_number)
            ws.cell(row=row, column=3, value=e.student.user.get_full_name().upper())
            # CA (col 4) and Exam (col 5) intentionally left blank

        for col, width in zip("ABCDE", [6, 22, 34, 8, 8]):
            ws.column_dimensions[col].width = width

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"{course.code}_{session_name.replace('/', '_')}_SCORE_SHEET.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class LecturerResultExcelUploadView(APIView):
    """
    POST /api/results/upload-excel/<allocation_pk>/  (multipart, field: file)
    Parses a filled score sheet: matches rows to approved enrollments by
    Registration No., score = CA + Exam. Creates/updates Results in the
    PENDING batch for the allocation's session/semester/programme.
    """
    permission_classes = [IsActiveLecturer]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, allocation_pk):
        try:
            allocation = CourseAllocation.objects.select_related(
                "course__programme", "session"
            ).get(pk=allocation_pk, lecturer=request.user.lecturer_profile)
        except CourseAllocation.DoesNotExist:
            return Response({"detail": "Allocation not found."}, status=404)

        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "No file provided."}, status=400)
        if not upload.name.lower().endswith(".xlsx"):
            return Response({"detail": "Only .xlsx score sheets are accepted."}, status=400)

        try:
            wb = load_workbook(upload, data_only=True)
        except Exception:
            return Response({"detail": "Could not read the Excel file. Is it a valid .xlsx?"}, status=400)
        ws = wb.active

        # Locate the column-header row ("#", "Registration No.", "Name", "CA", "Exam")
        header_row = None
        col_map = {}
        for row in ws.iter_rows(min_row=1, max_row=30):
            values = {str(c.value).strip().lower(): c.column for c in row if c.value is not None}
            if any("registration" in v for v in values):
                header_row = row[0].row
                for label, col in values.items():
                    if "registration" in label:
                        col_map["reg"] = col
                    elif label == "ca":
                        col_map["ca"] = col
                    elif label == "exam":
                        col_map["exam"] = col
                break
        if header_row is None or {"reg", "ca", "exam"} - set(col_map):
            return Response(
                {"detail": "Score sheet format not recognised. Use the downloaded template "
                           "(columns: #, Registration No., Name, CA, Exam)."},
                status=400,
            )

        enrollment_map = {
            e.student.matric_number.strip().upper(): e
            for e in Enrollment.objects.filter(
                allocation=allocation, status="APPROVED"
            ).select_related("student")
        }

        programme_type = allocation.course.programme.programme_type
        batch, _ = SemesterResultBatch.objects.get_or_create(
            session=allocation.session,
            semester=allocation.semester,
            programme_type=programme_type,
        )
        if batch.status == SemesterResultBatch.Status.APPROVED:
            return Response({"detail": "Results for this semester are already approved and locked."}, status=400)
        if batch.status == SemesterResultBatch.Status.REJECTED:
            # a corrected upload sends the batch back to the coordinator's queue
            batch.status = SemesterResultBatch.Status.PENDING
            batch.save()

        def component(value):
            """Blank cell -> None; otherwise Decimal (may raise InvalidOperation)."""
            if value is None or (isinstance(value, str) and not value.strip()):
                return None
            return Decimal(str(value).strip())

        uploaded, errors = 0, []
        seen = set()
        with transaction.atomic():
            for row in ws.iter_rows(min_row=header_row + 1):
                row_no = row[0].row
                reg = ws.cell(row=row_no, column=col_map["reg"]).value
                if reg is None or not str(reg).strip():
                    continue  # blank/footer row
                reg_key = str(reg).strip().upper()
                enrollment = enrollment_map.get(reg_key)
                if enrollment is None:
                    errors.append({"row": row_no, "registration_no": str(reg).strip(),
                                   "detail": "No approved enrollment for this registration number in this course."})
                    continue
                if reg_key in seen:
                    errors.append({"row": row_no, "registration_no": reg_key,
                                   "detail": "Duplicate row for this registration number."})
                    continue

                raw_ca = ws.cell(row=row_no, column=col_map["ca"]).value
                raw_exam = ws.cell(row=row_no, column=col_map["exam"]).value
                try:
                    ca = component(raw_ca)
                    exam = component(raw_exam)
                except InvalidOperation:
                    errors.append({"row": row_no, "registration_no": reg_key,
                                   "detail": f"CA/Exam must be numbers (got CA={raw_ca!r}, Exam={raw_exam!r})."})
                    continue
                if ca is None or exam is None:
                    missing = "CA and Exam are" if ca is None and exam is None else (
                        "CA is" if ca is None else "Exam is")
                    errors.append({"row": row_no, "registration_no": reg_key,
                                   "detail": f"{missing} empty — fill in both columns."})
                    continue
                if ca < 0 or exam < 0:
                    errors.append({"row": row_no, "registration_no": reg_key,
                                   "detail": f"CA and Exam cannot be negative (got CA={ca}, Exam={exam})."})
                    continue
                score = ca + exam
                if score > Decimal("100"):
                    errors.append({"row": row_no, "registration_no": reg_key,
                                   "detail": f"Total score {score} exceeds 100."})
                    continue

                Result.objects.update_or_create(
                    enrollment=enrollment,
                    defaults={"score": score, "batch": batch},
                )
                seen.add(reg_key)
                uploaded += 1

        missing = sorted(set(enrollment_map) - seen)
        return Response({
            "uploaded": uploaded,
            "errors": errors,
            "students_without_scores": missing,
            "batch": batch.id,
            "batch_status": batch.status,
        })


class BatchListView(APIView):
    """GET /api/results/batches/  — list all result batches."""
    permission_classes = [IsAdmin]

    def get(self, request):
        qs = SemesterResultBatch.objects.all().select_related("session", "approved_by")
        s = request.query_params.get("status")
        if s:
            qs = qs.filter(status=s)
        return Response(SemesterResultBatchSerializer(qs, many=True).data)

    def post(self, request):
        serializer = SemesterResultBatchSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=201)
        return Response(serializer.errors, status=400)


class ApproveBatchView(APIView):
    """POST /api/results/batches/<pk>/approve/  — PG Coordinator approves batch."""
    permission_classes = [IsPGCoordinator]

    def post(self, request, pk):
        try:
            batch = SemesterResultBatch.objects.get(pk=pk)
        except SemesterResultBatch.DoesNotExist:
            return Response({"detail": "Not found."}, status=404)

        action = request.data.get("action")
        comment = request.data.get("comment", "")
        if action == "approve":
            batch.status = SemesterResultBatch.Status.APPROVED
            batch.coordinator_comment = comment
            batch.approved_at = timezone.now()
            batch.approved_by = request.user
            batch.save()
            recompute_gpa_for_batch(batch)
            # Notify all affected students
            results = Result.objects.filter(batch=batch).select_related(
                "enrollment__student__user"
            )
            for r in results:
                notify(
                    r.enrollment.student.user,
                    "Results Published",
                    f"Your {batch.semester} semester {batch.session.name} results are now available.",
                    "RESULT",
                )
            return Response({"detail": "Batch approved and students notified."})
        elif action == "reject":
            batch.status = SemesterResultBatch.Status.REJECTED
            batch.coordinator_comment = comment
            batch.save()
            recompute_gpa_for_batch(batch)
            return Response({"detail": "Batch rejected."})
        return Response({"detail": "Invalid action."}, status=400)
